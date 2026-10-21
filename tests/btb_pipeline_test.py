"""Tests for the BTB pipeline in scout/supporting_data/btb.

Covers the unit conversions, the cost and source-note helpers in
populate_meas_in_from_btb.py, the crosswalk builder's matching helpers, a
consistency check on the committed tech_crosswalk.csv, and a small end-to-end
run of the populate step on synthetic data.
"""

import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

BTB_DIR = Path(__file__).parent.parent / "scout" / "supporting_data" / "btb"
# The BTB scripts import each other as top-level modules, and meas_pkg_gen.py
# is used to check that source citations are in the form it reads
sys.path.append(str(BTB_DIR))
sys.path.append(str(Path(__file__).parent.parent))

import build_tech_crosswalk as bt  # noqa: E402
import populate_meas_in_from_btb as pop  # noqa: E402
import unit_conversions as uc  # noqa: E402
from meas_pkg_gen import parse_source_details  # noqa: E402


# ==========================================================================
# unit_conversions
# ==========================================================================


@pytest.mark.parametrize("metric", [
    "SEER1", "SEER2", "EER", "CEER", "IEER", "HSPF"])
def test_btu_per_wh_metrics_divide_by_3412(metric):
    assert uc.convert(metric, "COP", 15) == pytest.approx(15 / 3.412)


def test_afue_is_percent_or_fraction():
    # Most BTB rows give AFUE as 0-100, some (oil boilers) as a 0-1 fraction
    assert uc.convert("AFUE", "AFUE", 90) == pytest.approx(0.9)
    assert uc.convert("AFUE", "COP", 0.86) == pytest.approx(0.86)


def test_ratio_and_passthrough_metrics():
    assert uc.convert("COP", "BTU out/BTU in", 3.9) == 3.9
    assert uc.convert("Heating COP", "COP", 3.6) == 3.6
    assert uc.convert("UEF", "UEF", 3.33) == 3.33
    assert uc.convert("Combined Energy Factor", "CEF", 11) == 11


@pytest.mark.parametrize("metric, unit, value", [
    ("Annual Energy Use", "COP", 5),   # no rule for the metric
    ("UEF", "COP", 5),                 # passthrough only to its own unit
    (None, "COP", 5),
    ("SEER1", None, 5),
    ("SEER1", "COP", "not a number"),
])
def test_convert_returns_none_without_a_rule(metric, unit, value):
    assert uc.convert(metric, unit, value) is None


def test_to_btb_metric_is_the_inverse_of_convert():
    for metric in ("SEER1", "IEER", "EER"):
        back = uc.to_btb_metric(metric, "COP", uc.convert(metric, "COP", 14.25))
        assert back == pytest.approx(14.25)
    assert uc.to_btb_metric("UEF", "UEF", 3.3) == 3.3


def test_to_btb_metric_afue_scale_follows_the_btb_row():
    # percent-scale row (typical 90) versus fraction-scale row (typical 0.86)
    assert uc.to_btb_metric("AFUE", "AFUE", 0.9, btb_scale_ref=90) == \
        pytest.approx(90)
    assert uc.to_btb_metric("AFUE", "AFUE", 0.86, btb_scale_ref=0.86) == \
        pytest.approx(0.86)
    assert uc.to_btb_metric("Annual Energy Use", "COP", 5) is None


# ==========================================================================
# populate_meas_in_from_btb: name, unit and cell helpers
# ==========================================================================


@pytest.mark.parametrize("name, tier", [
    ("(R) Ref. Case ASHP", "Ref. Case"),
    ("(R) Min. Efficiency ASHP FS (NG Furnace)", "Min. Efficiency"),
    ("(R) ESTAR HP FS (NG Furnace)", "ESTAR"),
    ("(R) ENERGY STAR Dishwasher", "ESTAR"),
    ("(R) Best Gas WH", "Best"),
    ("(R) Induction Range FS", "Best"),
    ("(R) Brk. HP FS (NG Furnace)", None),
    (np.nan, None),
])
def test_tier_for_name(name, tier):
    assert pop.tier_for_name(name) == tier


def test_sector_for_name():
    assert pop.sector_for_name("(R) Ref. Case ASHP") == "residential"
    assert pop.sector_for_name("(C) Ref. Case GSHP") == "commercial"
    assert pop.sector_for_name("Ref. Case ASHP") is None


@pytest.mark.parametrize("units, info", [
    ("2023$/unit", (2023, False)),
    ("2022$/unit", (2022, False)),
    ("2023$/kBtu/h heating", (2023, True)),
    ("2022$/kBtu/h water heating", (2022, True)),
    ("2023$/kBtu/h cooking", (2023, True)),
    ("2022$/kBtu/h refrigeration", None),   # not a BTB cost unit
    ("2022$/ft^2 floor", None),
    (np.nan, None),
])
def test_cost_units_info(units, info):
    assert pop.cost_units_info(units) == info


def test_relabel_cost_units():
    assert pop.relabel_cost_units("2022$/unit") == "2023$/unit"
    assert pop.relabel_cost_units("2022$/kBtu/h water heating") == \
        "2023$/kBtu/h water heating"


@pytest.mark.parametrize("cell, unkeyed", [
    (np.nan, True),
    (2.3, True),
    ("2.3", True),
    ("new: 100; existing: 120", True),
    ("ASHP: new: 100; ASHP: existing: 120", False),
    ("heating: 2.58; cooling: 4.4", False),
])
def test_is_unkeyed(cell, unkeyed):
    assert pop.is_unkeyed(cell) is unkeyed


def test_find_and_replace_value_for_key():
    cell = "furnace (NG): 0.9; central AC: 4.28; room AC: 3.52"
    assert pop.find_value_for_key(cell, "central AC") == "4.28"
    assert pop.find_value_for_key(cell, "boiler") is None
    assert pop.find_value_for_key("2.3", "anything") == "2.3"
    assert pop.find_value_for_key(np.nan, "anything") is None
    assert pop.replace_value_for_key(cell, "central AC", "4.4") == \
        "furnace (NG): 0.9; central AC: 4.4; room AC: 3.52"
    assert pop.replace_value_for_key("2.3", "anything", "4") == "4"
    assert pop.replace_value_for_key(np.nan, "anything", "4") == "4"


def test_replace_cost_for_key():
    cell = "ASHP: new: 100; ASHP: existing: 120; central AC: new: 5; " \
        "central AC: existing: 6"
    out = pop.replace_cost_for_key(cell, "ASHP", {"new": 111, "existing": 222})
    assert out == "ASHP: new: 111; ASHP: existing: 222; " \
        "central AC: new: 5; central AC: existing: 6"
    # a bare cost cell is rewritten whole
    assert pop.replace_cost_for_key(
        "new: 1; existing: 2", "ASHP", {"new": 7, "existing": 8}) == \
        "new: 7; existing: 8"


def test_owner_tech():
    assert pop.owner_tech({"Switched to Technology": "ASHP",
                           "Baseline Technology": "furnace (NG); central AC"}
                          ) == "ASHP"
    # switching to "same" leaves the baseline technology as the owner
    assert pop.owner_tech({"Switched to Technology": "same",
                           "Baseline Technology": "all"}) == "all"
    assert pop.owner_tech({"Baseline Technology": "furnace (NG); AC"}) is None
    assert pop.owner_tech({}) is None


def test_perf_key_for_tech():
    keyed = "furnace (NG): 0.9; central AC: 4.28"
    end_use = "heating: 2.58; cooling: 4.4"
    assert pop.perf_key_for_tech(keyed, "central AC", "SEER1", None) == \
        "central AC"
    # a cell keyed by end use is resolved from the BTB metric
    assert pop.perf_key_for_tech(end_use, "ASHP", "SEER1", None) == "cooling"
    assert pop.perf_key_for_tech(end_use, "ASHP", "HSPF", None) == "heating"
    assert pop.perf_key_for_tech(end_use, "ASHP", "Energy Use", None) is None
    # a bare value belongs to the owner only
    assert pop.perf_key_for_tech("2.3", "HPWH", "UEF", "HPWH") == "HPWH"
    assert pop.perf_key_for_tech("2.3", "other", "UEF", "HPWH") is None


def test_metric_conflicts_with_token():
    assert pop.metric_conflicts_with_token("comm_GSHP-cool", "Heating COP")
    assert pop.metric_conflicts_with_token("rooftop_ASHP-heat", "IEER")
    assert not pop.metric_conflicts_with_token("rooftop_ASHP-cool", "IEER")
    assert not pop.metric_conflicts_with_token("ASHP", "SEER1")


def test_parse_currency_and_zero_placeholder():
    assert pop.parse_currency("$5,158") == 5158
    assert pop.parse_currency(12.5) == 12.5
    assert pop.parse_currency(np.nan) is None
    assert pop.parse_currency("n/a") is None
    assert pop.is_zero_cost_placeholder("0")
    assert pop.is_zero_cost_placeholder(0)
    assert pop.is_zero_cost_placeholder("new: 0; existing: 0")
    assert not pop.is_zero_cost_placeholder("new: 5; existing: 6")
    assert not pop.is_zero_cost_placeholder(np.nan)


# ==========================================================================
# populate_meas_in_from_btb: BTB rows, capacity and cost
# ==========================================================================


def btb_row(**kwargs):
    """A BTB row with the columns the cost functions read."""

    values = {
        "Regression metric 1 - Metric": "Cooling Capacity",
        "Regression metric 1 - Unit": "Tons",
        "Regression metric 1 - Typical": 3.0,
        "Regression metric 2 - Metric": "SEER1",
        "Regression metric 2 - Typical": 15.0,
        "Regression metric 2 - Coefficient-Mid": 600.0,
        "Typical unit multiplier": np.nan,
        "Installation Multiplier - New Construction": 1.2,
        "Installation Multiplier - Retrofit": 1.5,
        "Typical New Construction Installed Cost ($2023) - Mid": "$10,000",
        "Typical Retrofit Installed Cost ($2023) - Mid": "$11,000",
    }
    values.update(kwargs)
    return pd.Series(values)


@pytest.mark.parametrize("metric, unit, typical, kbtuh", [
    ("Cooling Capacity", "Tons", 3.0, 36.0),
    ("Heating Capacity", "kBtu/h", 80.0, 80.0),
    ("Heating Capacity", "MBH", 250.0, 250.0),
    ("Heating Capacity", "Btu/hr", 100000.0, 100.0),
    ("Heating Capacity", "kW", 7.5, 7.5 * 3.412),
    ("Heat Output", "kBtu/h", 157.0, 157.0),
    ("Combined Idle Energy Rate", "BTU/h", 10350.0, 10.35),
    ("Combined Idle Energy Rate", "kW", 1.4, 1.4 * 3.412),
])
def test_capacity_kbtuh(metric, unit, typical, kbtuh):
    row = btb_row(**{"Regression metric 1 - Metric": metric,
                     "Regression metric 1 - Unit": unit,
                     "Regression metric 1 - Typical": typical})
    assert pop.capacity_kbtuh(row) == pytest.approx(kbtuh)


def test_capacity_kbtuh_unknown_or_override():
    volume = btb_row(**{"Regression metric 1 - Metric": "Nominal Volume",
                        "Regression metric 1 - Unit": "Gallons",
                        "Regression metric 1 - Typical": 100.0})
    assert pop.capacity_kbtuh(volume) is None
    # the commercial gas storage water heater uses the sheet's 199 kBtu/h
    volume["btb_sector"], volume["tech_id"] = "commercial", 65
    assert pop.capacity_kbtuh(volume) == 199.0
    airflow = btb_row(**{"Regression metric 1 - Metric": "Airflow",
                         "Regression metric 1 - Unit": "CFM"})
    assert pop.capacity_kbtuh(airflow) is None


def test_installation_multiplier_direct():
    row = btb_row()
    assert pop.installation_multiplier(row, "new") == 1.2
    assert pop.installation_multiplier(row, "retrofit") == 1.5


def pricing_row(retail, installed_new, installed_ret):
    """A commercial-style row: no multiplier columns, three price points."""

    row = {}
    for q, r, n, t in zip(("Low", "Mid", "High"), retail, installed_new,
                          installed_ret):
        row[f"Typical Retail Price ($2023) - {q}"] = r
        row[f"Typical New Construction Installed Cost ($2023) - {q}"] = n
        row[f"Typical Retrofit Installed Cost ($2023) - {q}"] = t
    return pd.Series(row)


def test_installation_multiplier_recovered_from_prices():
    # installed = retail + 50: an adder, so a multiplier of 1
    adder = pricing_row([100, 200, 300], [150, 250, 350], [160, 260, 360])
    assert pop.installation_multiplier(adder, "new") == 1.0
    # installed = retail * 1.2
    mult = pricing_row([100, 200, 300], [120, 240, 360], [150, 300, 450])
    assert pop.installation_multiplier(mult, "new") == pytest.approx(1.2)
    assert pop.installation_multiplier(mult, "retrofit") == pytest.approx(1.5)
    # no retail price: cannot be recovered
    missing = pricing_row([np.nan] * 3, [1, 2, 3], [1, 2, 3])
    assert pop.installation_multiplier(missing, "new") is None


def test_compute_cost_at_typical_without_a_performance():
    cost, used = pop.compute_cost(btb_row())
    assert cost == {"new": 10000, "existing": 11000}
    assert used is False


def test_compute_cost_moves_with_performance():
    # retail change = 600 * (16 - 15) = 600; installed moves by 1.2x and 1.5x
    cost, used = pop.compute_cost(btb_row(), metric_idx=2, perf_value=16.0)
    assert cost == {"new": 10720, "existing": 11900}
    assert used is True
    # at the typical performance the cost is BTB's own
    same, used = pop.compute_cost(btb_row(), metric_idx=2, perf_value=15.0)
    assert same == {"new": 10000, "existing": 11000} and used is True
    # a unit multiplier scales the retail change (e.g. per-ton regressions)
    scaled, _ = pop.compute_cost(
        btb_row(**{"Typical unit multiplier": 2.0}), 2, 16.0)
    assert scaled["new"] == 10000 + round(1.2 * 600 * 2)


def test_compute_cost_falls_back_when_the_regression_is_unusable():
    # metric 1 is a size, not a performance metric
    cost, used = pop.compute_cost(btb_row(), metric_idx=1, perf_value=16.0)
    assert cost == {"new": 10000, "existing": 11000} and used is False
    # no installed cost at all
    empty = btb_row(**{
        "Typical New Construction Installed Cost ($2023) - Mid": np.nan})
    assert pop.compute_cost(empty) == (None, False)


# ==========================================================================
# populate_meas_in_from_btb: composite costs
# ==========================================================================


def furnace_data(cost_new="$3,668", cost_ret="$3,861"):
    """BTB data and crosswalk with a residential gas and oil furnace."""

    rows, crosswalk = [], {}
    for token, tech_id, new, ret in (
            ("furnace (NG)", 91, cost_new, cost_ret),
            ("furnace (distillate)", 92, "$5,158", "$5,350")):
        rows.append({
            "tech_id": tech_id, "Projection Year": 2023,
            "Projection Scenario": "Reference",
            "Typical New Construction Installed Cost ($2023) - Mid": new,
            "Typical Retrofit Installed Cost ($2023) - Mid": ret})
        crosswalk[(token, "residential", "Ref. Case")] = pd.Series({
            "sector": "residential", "btb_technology_id": tech_id,
            "projection_year": 2023, "projection_scenario": "Reference"})
    return {"residential": pd.DataFrame(rows)}, crosswalk


def test_dual_fuel_from_name():
    assert pop.dual_fuel_from_name("(C) ESTAR HP Pt. FS (NG Furnace)") == "gas"
    assert pop.dual_fuel_from_name(
        "(C) ESTAR HP Pt. FS (Other Fossil Furnace)") == "oil"
    assert pop.dual_fuel_from_name("(C) ESTAR HP FS (NG Furnace)") is None
    assert pop.dual_fuel_from_name(np.nan) is None


def test_composite_adjustment_dual_fuel_adds_the_furnace_for_new_only():
    data, crosswalk = furnace_data()
    notes = "For new buildings, also add retail cost of typical gas furnace."
    extra, note = pop.composite_adjustment(
        notes, "residential", crosswalk, data)
    assert extra == {"new": 3668, "existing": 0}
    assert "gas furnace" in note
    # the oil furnace comes from the same wording
    oil, _ = pop.composite_adjustment(
        notes.replace("gas", "oil"), "residential", crosswalk, data)
    assert oil == {"new": 5158, "existing": 0}


def test_composite_adjustment_recognizes_dual_fuel_by_name():
    data, crosswalk = furnace_data()
    notes = "Uses residential system costs; existing buildings: tank removal"
    extra, _ = pop.composite_adjustment(
        notes, "residential", crosswalk, data,
        "(C) ESTAR HP Pt. FS (Other Fossil Furnace)")
    # the backup furnace is kept, so no tank removal for dual-fuel rows
    assert extra == {"new": 5158, "existing": 0}


def test_composite_adjustment_tank_removal_and_unhandled_notes():
    data, crosswalk = furnace_data()
    extra, _ = pop.composite_adjustment(
        "additional tank removal costs of 800 2014$", "residential",
        crosswalk, data, "(R) ESTAR HP FS (Other Fossil Furnace)")
    # 800 2014$ updated to 2023$ in the existing cost notes
    assert extra == {"new": 0, "existing": 1030}
    assert pop.composite_adjustment(
        "assume a breakthrough HP costs half the cost of the best", "r",
        crosswalk, data) is None
    assert pop.composite_adjustment(
        "plain note", "residential", crosswalk, data) == (
            {"new": 0, "existing": 0}, "")
    assert pop.composite_adjustment(
        np.nan, "residential", crosswalk, data) == (
            {"new": 0, "existing": 0}, "")
    # a dual-fuel row whose furnace is not in the crosswalk cannot be rebuilt
    assert pop.composite_adjustment(
        "also add retail cost of typical gas furnace", "residential",
        {}, data) is None


# ==========================================================================
# populate_meas_in_from_btb: sources and measure edits
# ==========================================================================


def sheet_for(columns):
    sheet = openpyxl.Workbook().active
    return sheet, {name: i + 1 for i, name in enumerate(columns)}


def test_btb_source_details_are_in_the_form_the_generator_reads():
    parsed = parse_source_details(pop.BTB_SOURCE_DETAILS)
    assert parsed["year"] == 2024
    assert parsed["pages"] is None
    assert parsed["url"] == "https://data.openei.org/submissions/8342"
    assert "Buildings Technology Baseline" in parsed["title"]


def test_update_sources_replaces_when_every_entry_came_from_btb():
    row = pd.Series({
        "Name": "(R) Best HPWH FS", "Baseline Technology": "all",
        "Switched to Technology": "HPWH",
        "Cost Source Notes": "EIA cost", "Cost Source Details": "EIA; x"})
    sheet, cols = sheet_for(["Cost Source Notes", "Cost Source Details"])
    diff = []
    pop.update_sources(
        row, "Cost", {"HPWH": "BTB installed cost"}, "new: 1; existing: 2",
        sheet, 2, cols, diff)
    assert sheet.cell(2, 1).value == "BTB installed cost"
    assert sheet.cell(2, 2).value == pop.BTB_SOURCE_DETAILS
    assert [d["column"] for d in diff] == [
        "Cost Source Notes", "Cost Source Details"]


def test_update_sources_adds_to_the_source_when_only_some_entries_changed():
    cell = ("furnace (NG): new: 1; furnace (NG): existing: 2; "
            "central AC: new: 3; central AC: existing: 4")
    row = pd.Series({
        "Name": "(R) Ref. Case NG Furnace & AC",
        "Baseline Technology": "furnace (NG); central AC",
        "Switched to Technology": np.nan,
        "Cost Source Notes": "EIA cost", "Cost Source Details": "EIA; x"})
    sheet, cols = sheet_for(["Cost Source Notes", "Cost Source Details"])
    diff = []
    pop.update_sources(
        row, "Cost", {"furnace (NG)": "BTB installed cost"}, cell,
        sheet, 2, cols, diff)
    notes, details = sheet.cell(2, 1).value, sheet.cell(2, 2).value
    assert notes.startswith("EIA cost; BTB installed cost")
    assert "(furnace (NG))" in notes
    assert details == "EIA; x\n" + pop.BTB_SOURCE_DETAILS
    # the added citation is not repeated if BTB is already cited
    row["Cost Source Details"] = details
    sheet2, cols2 = sheet_for(["Cost Source Notes", "Cost Source Details"])
    diff2 = []
    pop.update_sources(
        row, "Cost", {"furnace (NG)": "BTB installed cost"}, cell,
        sheet2, 2, cols2, diff2)
    assert sheet2.cell(2, 2).value is None
    assert [d["column"] for d in diff2] == ["Cost Source Notes"]


def test_update_sources_does_nothing_without_updates():
    row = pd.Series({"Name": "x"})
    sheet, cols = sheet_for(["Cost Source Notes", "Cost Source Details"])
    diff = []
    pop.update_sources(row, "Cost", {}, "0", sheet, 2, cols, diff)
    assert diff == [] and sheet.cell(2, 1).value is None


def test_apply_measure_edits_renames_and_logs():
    name = "(R) Best 120V HP Dryer FS"
    row = pd.Series({
        "Name": name, "Description": "Fuel switch from gas to 120VHP dryer",
        "Electrical Upgrade Costs": "panel: false; 240V circuit: false"})
    sheet, cols = sheet_for(["Name", "Description", "Electrical Upgrade Costs"])
    diff = []
    out = pop.apply_measure_edits(
        row, pop.MEASURE_EDITS[name], sheet, 2, cols, diff)
    assert out["Name"] == "(R) Best HP Dryer FS"
    assert out["Description"] == "Fuel switch from gas to HP dryer"
    assert out["Electrical Upgrade Costs"] == "panel: false; 240V circuit: true"
    assert {d["column"] for d in diff} == {
        "Name", "Description", "Electrical Upgrade Costs"}
    assert sheet.cell(2, 1).value == "(R) Best HP Dryer FS"
    # the original row is not modified
    assert row["Name"] == name


def test_apply_measure_edits_accepts_functions_of_the_old_value():
    row = pd.Series({
        "Name": "(C) Ref. Case GSHP",
        "Energy Performance": "comm_GSHP-heat: 3.2; comm_GSHP-cool: 4.13",
        "Performance Source Notes": "Typical efficiency level in 2023."})
    sheet, cols = sheet_for(["Name", "Energy Performance",
                             "Performance Source Notes"])
    out = pop.apply_measure_edits(
        row, pop.MEASURE_EDITS["(C) Ref. Case GSHP"], sheet, 2, cols, [])
    assert out["Energy Performance"] == \
        "comm_GSHP-heat: 3.2; comm_GSHP-cool: 6.33"
    assert out["Performance Source Notes"].startswith(
        "Typical efficiency level in 2023; cooling COP")


def test_refrigeration_edit_covers_all_eight_technologies():
    edit = pop.MEASURE_EDITS["(C) Ref. Case Refrigeration"]
    techs = set(pop.EIA_REFRIGERATION_COST)
    assert techs == set(pop.EIA_REFRIGERATION_LIFETIME)
    assert len(techs) == 8
    assert edit["Measure Type"] == "full service"
    for tech in techs:
        assert f"{tech}: new:" in edit["Installed Cost"]
        assert f"{tech}: existing:" in edit["Installed Cost"]
    # new (retail) is never above existing (installed)
    assert all(new <= existing
               for new, existing in pop.EIA_REFRIGERATION_COST.values())


def test_residential_ashp_cooling_by_tier():
    meas = pd.DataFrame({
        "Name": ["(R) Min. Efficiency ASHP FS (NG Furnace)",
                 "(R) ESTAR HP FS (NG Furnace)",
                 "(R) Best HP FS (NG Furnace)",
                 "(C) Best HP FS (NG Furnace & AC)",
                 "(R) Ref. Case GSHP"],
        "Switched to Technology": ["ASHP", "ASHP", "ASHP", "ASHP", "GSHP"],
        "Energy Performance": [
            "heating: 2.58; cooling: 4.4",
            "cold climates:heating:2.79; cold climates:cooling:4.69",
            "cooling:6.6; heating:3.6",
            "heating: 1; cooling: 9",
            "heating: 3.5; cooling: 5.07"]})
    assert pop.residential_ashp_cooling(meas) == {
        "Min. Efficiency": 4.4, "ESTAR": 4.69, "Best": 6.6}


# ==========================================================================
# build_tech_crosswalk
# ==========================================================================


def test_find_candidates_applies_include_and_exclude_keywords():
    df = pd.DataFrame({
        "tech_id": [1, 2, 3, 4, 4, np.nan],
        "search_text": [
            "air source heat pump | heating",
            "air source heat pump with new circuit",
            "ground source heat pump",
            "furnace gas-fired",
            "furnace gas-fired",
            "air source heat pump"]})
    found = bt.find_candidates(df, ["air source heat pump"], ["new circuit"])
    assert list(found["tech_id"]) == [1]
    # candidates are de-duplicated by technology and need an id
    assert list(bt.find_candidates(df, ["furnace"], [])["tech_id"]) == [4]


def test_build_tier_for_name_matches_populate():
    for name in ("(R) Ref. Case ASHP", "(R) ESTAR HP", "(R) Best Gas WH",
                 "(R) Min. Efficiency ASHP", "(R) Induction Range"):
        expected = pop.tier_for_name(name)
        assert bt.tier_for_name(name) in (expected, None)
    assert bt.tier_for_name("(R) Ref. Case Gas WH") == "Ref. Case"


def test_match_rules_are_well_formed():
    for token, rules in bt.MATCH_RULES.items():
        assert rules, token
        for sector, include, exclude in rules:
            assert sector in ("residential", "commercial"), token
            assert include and all(k == k.lower() for k in include + exclude)
    # refrigeration uses EIA entirely, so has no BTB rules
    assert not any("Commercial" in token and "Refriger" in token
                   or "Reach-In" in token or "Walk-In" in token
                   for token in bt.MATCH_RULES)


def test_cross_sector_tokens_have_a_rule_in_the_other_sector():
    for token in pop.CROSS_SECTOR_TECHS:
        sectors = {rule[0] for rule in bt.MATCH_RULES[token]}
        assert sectors == {"residential"}, token


# ==========================================================================
# the committed tech_crosswalk.csv
# ==========================================================================


@pytest.fixture(scope="module")
def crosswalk():
    return pd.read_csv(BTB_DIR / "tech_crosswalk.csv")


def test_crosswalk_columns_and_values(crosswalk):
    assert list(crosswalk.columns) == bt.CROSSWALK_COLUMNS
    assert not crosswalk.duplicated(
        ["scout_technology", "sector", "tier"]).any()
    assert set(crosswalk["sector"]) <= {"residential", "commercial"}
    assert set(crosswalk["tier"]) <= set(bt.TIER_DEFAULTS)
    assert set(crosswalk["bound"]) <= {"Typical", "High", "Existing"}
    assert not crosswalk["needs_review"].astype(bool).any()
    assert crosswalk["btb_technology_id"].notna().all()


def test_crosswalk_tiers_follow_the_defaults(crosswalk):
    tokens = crosswalk[~crosswalk["scout_technology"].str.startswith("(")]
    for (token, sector), group in tokens.groupby(
            ["scout_technology", "sector"]):
        assert set(group["tier"]) == set(bt.TIER_DEFAULTS), (token, sector)
        for _, row in group.iterrows():
            assert row["bound"] == bt.TIER_DEFAULTS[row["tier"]]["bound"]


def test_measure_level_rows_name_their_own_tier(crosswalk):
    measures = crosswalk[crosswalk["scout_technology"].str.startswith("(")]
    assert len(measures) == len(bt.MEASURE_RULES)
    for _, row in measures.iterrows():
        assert row["tier"] == bt.tier_for_name(row["scout_technology"])
        assert row["scout_technology"] in bt.MEASURE_RULES


def test_overrides_in_the_crosswalk_match_the_build_tables(crosswalk):
    for token, value in bt.COST_AT.items():
        assert set(crosswalk[crosswalk["scout_technology"] == token][
            "cost_at"]) == {value}
    for token, value in bt.KEEP_LIFETIME.items():
        assert set(crosswalk[crosswalk["scout_technology"] == token][
            "keep_lifetime"]) == {value}
    for token, value in bt.LIFETIME_VALUE.items():
        assert set(crosswalk[crosswalk["scout_technology"] == token][
            "lifetime_value"]) == {value}


# ==========================================================================
# end to end on synthetic data
# ==========================================================================


MEAS_COLUMNS = [
    "Name", "Description", "Measure Type", "Baseline Technology",
    "Switched to Technology", "Energy Performance", "Performance Units",
    "Performance Source Notes", "Performance Source Details",
    "Installed Cost", "Cost Units", "Cost Source Notes",
    "Cost Source Details", "Lifetime", "Lifetime Source Notes",
    "Lifetime Source Details"]


def make_btb_files(raw_dir):
    """Minimal BTB residential/commercial CSVs with one heat pump water
    heater (residential ID 169)."""

    raw_dir.mkdir()
    hpwh = {
        "Technology ID": 169, "Display Name": "Water Heater HP Tank",
        "Projection Year": 2023, "Projection Scenario": "Reference",
        "Regression metric 1 - Metric": "Nominal Volume",
        "Regression metric 1 - Unit": "Gallons",
        "Regression metric 1 - Typical": 60.0,
        "Regression metric 2 - Metric": "UEF",
        "Regression metric 2 - Unit": "Unitless",
        "Regression metric 2 - Typical": 3.33,
        "Regression metric 2 - High": 3.73,
        "Regression metric 2 - Coefficient-Mid": 500.0,
        "Installation Multiplier - New Construction": 1.0,
        "Installation Multiplier - Retrofit": 1.0,
        "Typical New Construction Installed Cost ($2023) - Mid": "$2,670",
        "Typical Retrofit Installed Cost ($2023) - Mid": "$3,170",
        "Lifetime (Years)": 15.1}
    pd.DataFrame([hpwh]).to_csv(raw_dir / "btb_residential.csv", index=False)
    commercial = {
        "Technology/Measure ID": 1, "Display Name": "unused",
        "Year": 2023, "Scenario": "Reference"}
    pd.DataFrame([commercial]).to_csv(
        raw_dir / "btb_commercial.csv", index=False)


def make_crosswalk_file(path):
    rows = []
    for tier, bound in (("Ref. Case", "Typical"), ("Best", "High")):
        rows.append({
            "scout_technology": "HPWH", "sector": "residential",
            "tier": tier, "btb_technology_id": 169,
            "btb_display_name": "Water Heater HP Tank",
            "btb_metric_name": "UEF", "projection_scenario": "Reference",
            "projection_year": 2023, "bound": bound, "needs_review": False,
            "notes": "", "cost_at": "", "measure_type": "",
            "keep_lifetime": "", "lifetime_value": ""})
    pd.DataFrame(rows).to_csv(path, index=False)


def make_meas_in_file(path):
    base = {c: None for c in MEAS_COLUMNS}
    best = {**base, "Name": "(R) Best HPWH FS", "Measure Type": "full service",
            "Baseline Technology": "all", "Switched to Technology": "HPWH",
            "Energy Performance": "3.0", "Performance Units": "UEF",
            "Performance Source Notes": "EIA value",
            "Performance Source Details": "EIA; U.S. EIA; 2023; 65; url",
            "Installed Cost": "new: 2000; existing: 2500",
            "Cost Units": "2022$/unit", "Cost Source Notes": "EIA cost",
            "Cost Source Details": "EIA; U.S. EIA; 2023; 65; url",
            "Lifetime": "12", "Lifetime Source Notes": "EIA life",
            "Lifetime Source Details": "EIA; U.S. EIA; 2023; 65; url"}
    breakthrough = {**best, "Name": "(R) Brk. HPWH FS"}
    no_tier = {**best, "Name": "(R) Some HPWH FS"}
    sheet = openpyxl.Workbook()
    ws = sheet.active
    ws.title = "meas_in"
    ws.append(MEAS_COLUMNS)
    for row in (best, breakthrough, no_tier):
        ws.append([row[c] for c in MEAS_COLUMNS])
    sheet.save(path)


@pytest.fixture
def pipeline_output(tmp_path, monkeypatch):
    make_btb_files(tmp_path / "raw")
    make_crosswalk_file(tmp_path / "crosswalk.csv")
    make_meas_in_file(tmp_path / "meas.xlsx")
    monkeypatch.setattr(pop, "RAW_DIR", tmp_path / "raw")
    monkeypatch.setattr(pop, "CROSSWALK_PATH", tmp_path / "crosswalk.csv")
    monkeypatch.setattr(pop, "XLSX_PATH", tmp_path / "meas.xlsx")
    monkeypatch.setattr(pop, "OUT_PATH", tmp_path / "out.xlsx")
    pop.main()
    return openpyxl.load_workbook(tmp_path / "out.xlsx")


def cells(ws, name):
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            return dict(zip(header, [c.value for c in row]))
    return None


def test_pipeline_updates_the_best_tier_from_btb(pipeline_output):
    row = cells(pipeline_output["meas_in"], "(R) Best HPWH FS")
    # Best takes BTB's High UEF; cost moves by 500 * (3.73 - 3.33) = 200
    assert float(row["Energy Performance"]) == pytest.approx(3.73)
    assert row["Installed Cost"] == "new: 2870; existing: 3370"
    assert float(row["Lifetime"]) == pytest.approx(15.1)
    # the single-cost 2022$ row is relabeled to BTB's dollar year
    assert row["Cost Units"] == "2023$/unit"


def test_pipeline_cites_btb_for_what_it_changed(pipeline_output):
    row = cells(pipeline_output["meas_in"], "(R) Best HPWH FS")
    for category in ("Performance", "Cost", "Lifetime"):
        assert row[f"{category} Source Details"] == pop.BTB_SOURCE_DETAILS
        assert "BTB" in row[f"{category} Source Notes"]
    assert "high performance" in row["Cost Source Notes"]


def test_pipeline_leaves_breakthrough_and_untiered_measures_alone(
        pipeline_output):
    ws = pipeline_output["meas_in"]
    for name in ("(R) Brk. HPWH FS", "(R) Some HPWH FS"):
        row = cells(ws, name)
        assert row["Energy Performance"] == "3.0"
        assert row["Installed Cost"] == "new: 2000; existing: 2500"
        assert row["Lifetime"] == "12"
        assert row["Cost Source Notes"] == "EIA cost"


def test_pipeline_logs_every_change_in_the_diff_sheet(pipeline_output):
    diff = pd.DataFrame(pipeline_output["BTB Diff"].values)
    diff.columns = diff.iloc[0]
    diff = diff.iloc[1:]
    assert set(diff["Name"]) == {"(R) Best HPWH FS"}
    assert {"Energy Performance", "Installed Cost", "Lifetime", "Cost Units",
            "Cost Source Notes", "Cost Source Details"} <= set(diff["column"])
    cost = diff[diff["column"] == "Installed Cost"].iloc[0]
    assert cost["old_value"] == "new: 2000; existing: 2500"
    assert cost["new_value"] == "new: 2870; existing: 3370"
