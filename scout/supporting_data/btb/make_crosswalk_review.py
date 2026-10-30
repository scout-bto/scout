"""Write tech_crosswalk_review.xlsx, a one-row-per-technology view of
tech_crosswalk.csv for review.

tech_crosswalk.csv repeats every (technology, sector) pair for the four
efficiency tiers, with the same BTB match each time. This collapses the
tiers so a reviewer can concentrate on the columns that matter:
scout_technology, sector, btb_technology_id, btb_display_name and
btb_performance_metric. Helper columns show how many meas_in rows use each
technology (in the entry's sector, and in the other sector, where the row
is skipped unless the technology is crosswalked across sectors) and any
known question about the mapping. The last three columns are left empty for
the reviewer.

Re-run this after editing tech_crosswalk.csv or FLAGS, from this directory:
    python make_crosswalk_review.py
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
CROSSWALK_PATH = BASE_DIR / "tech_crosswalk.csv"
XLSX_PATH = BASE_DIR / "bss_meas_v2.xlsx"
OUT_PATH = BASE_DIR / "tech_crosswalk_review.xlsx"

# (scout_technology, sector) -> open question about the mapping. Edit as
# questions are answered; the yellow-highlighted column in the output.
FLAGS = {
}

# Review rows for decisions that are not (yet) crosswalk entries: rows that
# are left as they are today and a proposed mapping for them. Shown with
# "(proposed)" in scout_technology and appended to the review table.
EXTRA_ROWS = []

COLUMN_WIDTHS = [26, 12, 10, 38, 38, 22, 10, 12, 36, 70, 14, 22, 40]


def sector_of(name):
    """Sector from a meas_in Name's (R)/(C) prefix."""

    return "residential" if str(name).startswith("(R)") else "commercial"


def usage_counts():
    """Number of meas_in rows using each technology token, by sector."""

    meas = pd.read_excel(XLSX_PATH, sheet_name="meas_in")
    counts = {}
    for _, row in meas.iterrows():
        sector = sector_of(row["Name"])
        seen = set()
        for col in ["Baseline Technology", "Switched to Technology"]:
            if isinstance(row[col], str):
                for token in row[col].split(";"):
                    token = token.strip()
                    if token and token not in seen:
                        seen.add(token)
                        key = (token, sector)
                        counts[key] = counts.get(key, 0) + 1
        counts[(row["Name"], sector)] = 1
    first_example = meas.drop_duplicates("Name").set_index("Name")
    return counts, first_example


def build_review():
    """The review table, one row per (technology, sector) crosswalk pair."""

    crosswalk = pd.read_csv(CROSSWALK_PATH)
    pairs = crosswalk.drop_duplicates(
        ["scout_technology", "sector"]).copy()
    counts, _ = usage_counts()

    rows = []
    for _, pair in pairs.iterrows():
        tech, sector = pair["scout_technology"], pair["sector"]
        other = "commercial" if sector == "residential" else "residential"
        rows.append({
            "scout_technology": tech, "sector": sector,
            "btb_technology_id": int(pair["btb_technology_id"]),
            "btb_display_name": pair["btb_display_name"],
            "btb_mapping_name": pair["btb_mapping_name"]
            if isinstance(pair["btb_mapping_name"], str) else "",
            "btb_performance_metric": pair["btb_metric_name"],
            "n_rows_same_sector": counts.get((tech, sector), 0),
            "n_rows_other_sector": counts.get((tech, other), 0),
            "measure_type_set": pair.get("measure_type")
            if isinstance(pair.get("measure_type"), str) else "",
            "flag_for_review": FLAGS.get((tech, sector), ""),
            "mapping_ok (Y/N)": "", "correct_btb_id_or_name": "",
            "colleague_notes": "",
        })
    ids = crosswalk.drop_duplicates("scout_technology").set_index(
        "scout_technology")
    for extra in EXTRA_ROWS:
        row = {"mapping_ok (Y/N)": "", "correct_btb_id_or_name": "",
               "colleague_notes": "", **extra}
        mapping = ids["btb_mapping_name"].get("rooftop_ASHP-heat")
        row["btb_mapping_name"] = mapping if isinstance(mapping, str) else ""
        rows.append(row)
    review = pd.DataFrame(rows)
    return review.sort_values(["sector", "scout_technology"]).reset_index(
        drop=True)


def write_workbook(review):
    """Write the review table and an About sheet, formatted for reading."""

    about = pd.DataFrame({"About this file": [
        "One row per (Scout technology, sector) in tech_crosswalk.csv; each "
        "repeats for 4 tiers there.",
        "The tier columns are the same for every technology: Ref. Case = "
        "BTB Typical, Best = BTB High, Min. Efficiency and ESTAR = keep the "
        "performance already in meas_in (cost from BTB regression at that "
        "performance). All use 2023 Reference.",
        "Please check: scout_technology -> btb_display_name (and "
        "btb_performance_metric). Fill in the last three columns where a "
        "row is wrong or needs a comment.",
        "n_rows_same_sector = meas_in rows in the entry's sector that use "
        "the technology. n_rows_other_sector = rows in the other sector "
        "that use the same token; those rows are skipped unless the token "
        "is crosswalked across sectors on purpose.",
        "Rows whose scout_technology starts with (R) or (C) are "
        "measure-level rows for measures with no technology token "
        "(baseline 'all'); measure_type_set shows a Measure Type change.",
        "Rows with (proposed) in scout_technology are not in "
        "tech_crosswalk.csv: they are rows left as they are today, with a "
        "proposed mapping to review.",
        "Technologies with no mapping are in "
        "technologies_without_btb_rule.xlsx.",
    ]})
    with pd.ExcelWriter(OUT_PATH, engine="openpyxl") as writer:
        review.to_excel(writer, sheet_name="Crosswalk review", index=False)
        about.to_excel(writer, sheet_name="About", index=False)

    workbook = load_workbook(OUT_PATH)
    sheet = workbook["Crosswalk review"]
    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")
    flag_col = list(review.columns).index("flag_for_review")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row[flag_col].value:
            row[flag_col].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    about_sheet = workbook["About"]
    about_sheet.column_dimensions["A"].width = 120
    for row in about_sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    about_sheet["A1"].font = Font(bold=True)
    workbook.save(OUT_PATH)


if __name__ == "__main__":
    table = build_review()
    write_workbook(table)
    print(f"Wrote {len(table)} rows to {OUT_PATH} "
          f"({int((table.flag_for_review != '').sum())} flagged).")
