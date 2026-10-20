"""Propose BTB-derived cost/performance/lifetime values for meas_in.

Reads bss_meas_v2.xlsx (sheet "meas_in"), tech_crosswalk.csv (see
build_tech_crosswalk.py), and the raw BTB CSVs, and for every (measure row,
technology) pair covered by a non-`needs_review` crosswalk entry, computes
what the Energy Performance / Installed Cost / Lifetime cells for that
technology *would* be if sourced from BTB.

This never modifies bss_meas_v2.xlsx. It writes bss_meas_v2_btb_proposed.xlsx
alongside it: a copy of meas_in with proposed values substituted in (only
for the specific technology substrings that were resolved -- everything
else in a shared cell, e.g. other technologies not covered by the
crosswalk, is left byte-for-byte as-is), plus a "BTB Diff" sheet listing
every change with its old/new value and matched BTB source, for manual
review before merging anything back into the curated original.

Usage (from this directory):
    python populate_meas_in_from_btb.py
"""

import re
from pathlib import Path

import openpyxl
import pandas as pd

from unit_conversions import convert, to_btb_metric

BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "raw"
XLSX_PATH = BASE_DIR / "bss_meas_v2.xlsx"
CROSSWALK_PATH = BASE_DIR / "tech_crosswalk.csv"
OUT_PATH = BASE_DIR / "bss_meas_v2_btb_proposed.xlsx"

SHEET_NAME = "meas_in"

TIER_KEYWORDS = [
    # Longer/more specific strings first so e.g. "ESTAR" doesn't also match
    # a "Min. Efficiency ESTAR ..." style name meant for a different tier.
    ("Min. Efficiency", "Min. Efficiency"),
    ("ESTAR", "ESTAR"),
    ("ENERGY STAR", "ESTAR"),
    ("Best", "Best"),
    ("Ref. Case", "Ref. Case"),
    # Induction ranges are the high-efficiency range option (the existing
    # cost notes call them "High").
    ("Induction", "Best"),
]

# Crosswalk "bound" -> BTB regression-metric column suffix, for tiers whose
# performance comes from BTB. The special bound "Existing" (Min. Efficiency
# and ESTAR) instead keeps the performance already in meas_in, which is set
# by a standard/ENERGY STAR spec rather than by BTB.
REGRESSION_BOUND_COL = {"Typical": "Typical", "High": "High"}
EXISTING_BOUND = "Existing"

# Technologies deliberately crosswalked to the other sector's BTB data
# (commercial measures using residential-type equipment), exempt from the
# row-sector check in main().
CROSS_SECTOR_TECHS = {
    "res_type_central_AC", "res_type_gasHP-cool", "res_type_gasHP-heat",
    "electric_res-heat", "wall-window_room_AC"}

# Commercial rows whose cost notes say they use residential system costs
# (divided by the typical residential capacity, 36 kBtu/h) take the
# residential ASHP entry for their cost only; performance and lifetime stay
# as they are.
_RESIDENTIAL_SYSTEM_NOTE = re.compile(r"residential system", re.IGNORECASE)

# Technologies covered by build_tech_crosswalk.py that participate in a
# shared "heating"/"cooling"/"ventilation" Performance Units key rather
# than their own tech-name key (seen in combined-package meas_in rows,
# e.g. "(C) Ref. Case NG Boiler & Chiller"). Used only as a fallback when a
# row's Performance Units cell has no key matching the technology name
# directly.
END_USE_FALLBACK = {
    "heating": [
        "gas_boiler", "elec_boiler", "oil_boiler", "gas_furnace",
        "oil_furnace", "elec_res-heater", "rooftop_ASHP-heat",
        "comm_GSHP-heat", "pkg_terminal_HP-heat", "resistance heat",
        "furnace (NG)", "furnace (distillate)", "boiler (distillate)",
        "ASHP", "GSHP", "HPWH"],
    "cooling": [
        "gas_chiller", "centrifugal_chiller", "scroll_chiller",
        "screw_chiller", "reciprocating_chiller", "rooftop_AC",
        "rooftop_ASHP-cool", "comm_GSHP-cool", "pkg_terminal_AC-cool",
        "pkg_terminal_HP-cool", "gas_eng-driven_RTAC", "central AC",
        "room AC", "ASHP", "GSHP"],
    "ventilation": ["VAV_Vent", "CAV_Vent"],
}


def normalize_btb(df, sector):
    """Rename sector-specific columns to a common schema."""

    id_col = "Technology ID" if "Technology ID" in df.columns \
        else "Technology/Measure ID"
    rename = {id_col: "tech_id"}
    if sector == "commercial":
        rename.update({
            "Year": "Projection Year", "Scenario": "Projection Scenario"})
    df = df.rename(columns=rename)
    df["btb_sector"] = sector
    return df


def load_btb_data():
    """Load and normalize both BTB CSVs, keyed by sector name."""

    out = {}
    for sector, fname in [
            ("residential", "btb_residential.csv"),
            ("commercial", "btb_commercial.csv")]:
        path = RAW_DIR / fname
        if not path.exists():
            raise FileNotFoundError(
                f"{path} not found -- run download_btb_data.py first.")
        out[sector] = normalize_btb(
            pd.read_csv(path, low_memory=False), sector)
    return out


def tier_for_name(name):
    """Infer a meas_in efficiency tier from a measure's Name string."""

    if not isinstance(name, str):
        return None
    for keyword, tier in TIER_KEYWORDS:
        if keyword in name:
            return tier
    return None


# Cost Units a BTB installed cost can be written into: whole-unit costs
# ("2023$/unit") or costs per kBtu/h of heating, cooling or water heating
# capacity, in 2023$ (BTB's dollar year) or 2022$. A 2022$ row is only
# replaced when it holds a single cost (nothing in the row stays in the old
# dollar year), and its Cost Units are then relabeled as 2023$.
# Per-capacity costs are the whole-unit $2023 cost divided by the BTB
# regression's Typical capacity expressed in kBtu/h (see capacity_kbtuh),
# as done in the "BTB Key Costs" sheet's "div capacity" columns.
_COST_UNITS = re.compile(
    r"(20\d\d)\$/(unit|kBtu/h (?:heating|cooling|water heating|cooking))$")
COST_UNITS_YEARS = {2022, 2023}


def cost_units_info(units):
    """Return (dollar year, is_per_capacity) for a supported Cost Units
    string, or None if the units are not supported."""

    match = _COST_UNITS.match(units) if isinstance(units, str) else None
    if not match or int(match.group(1)) not in COST_UNITS_YEARS:
        return None
    return int(match.group(1)), match.group(2) != "unit"


def relabel_cost_units(units):
    """The same Cost Units, in 2023$."""

    return units.replace(units[:4], "2023", 1)


# Capacity in kBtu/h of the typical BTB commercial unit where BTB describes
# it by volume rather than capacity. Keyed by (sector, BTB technology ID).
BTB_UNIT_CAPACITY_KBTUH = {
    # Commercial gas storage water heater (BTB: 100 gallons). The capacity
    # is the 199 kBtu/h that the BTB Key Costs sheet divides its cost by
    # (cell AG55), close to the 200 kBtu/h typical EIA commercial gas water
    # heater capacity Scout assumes (ecm_cost_convert.json).
    ("commercial", 65): 199.0,
}

# BTB capacity unit (lowercased) -> multiplier to kBtu/h.
CAPACITY_TO_KBTUH = {
    "tons": 12.0, "kbtu/h": 1.0, "mbh": 1.0, "btu/h": 0.001,
    "btu/hr": 0.001, "kw": 3.412}

# Breakthrough ("Brk.") measures are assumed costs, not taken from BTB.
_BREAKTHROUGH_NAME = re.compile(r"\bBrk\.?\b", re.IGNORECASE)

# Cost Source Notes phrases describing a cost built from more than the
# technology's own BTB cost. These are rebuilt from BTB pieces where
# possible (see composite_adjustment).
_ALSO_ADD_FURNACE = re.compile(
    r"also add retail cost of typical (gas|oil) furnace", re.IGNORECASE)
_TANK_REMOVAL = re.compile(r"tank removal", re.IGNORECASE)
_ELEC_BOILER_COST = re.compile(
    r"switched to elec boiler|commercial electric boiler unit",
    re.IGNORECASE)
_UNHANDLED_COMPOSITE = re.compile(
    r"half the cost|halve|add retail|also add", re.IGNORECASE)

# Oil-tank removal added to existing-construction costs when distillate or
# propane equipment is replaced (800 2014$ = 1030 2023$ per the 2015 EIA
# report cited in the existing cost notes). Not a BTB value.
TANK_REMOVAL_COST = 1030

# Technologies whose cost entry, in the rows described by _ELEC_BOILER_COST,
# holds the cost of the electric boiler switched to.
BOILER_COST_KEYS = {"gas_boiler", "oil_boiler", "boiler (NG)",
                    "boiler (distillate)"}

# Commercial rows representing a residential-type furnace in a commercial
# setting say so in their cost notes; the commercial furnace tokens then take
# the residential furnace's BTB entry (cost, performance and lifetime).
_RESIDENTIAL_FURNACE_NOTE = re.compile(r"residential furnaces", re.IGNORECASE)
RESIDENTIAL_FURNACE_ALIAS = {
    "oil_furnace": "furnace (distillate)", "gas_furnace": "furnace (NG)"}

# EIA typical-case costs (2022$/kBtu/h cooling capacity) for the commercial
# refrigeration technologies, from the "Updated Buildings Sector Appliance and
# Equipment Costs and Efficiencies" report, Appendix C: new = Retail Equipment
# Cost and existing = Total Installed Cost (2022 Typical column) divided by
# the typical unit's cooling capacity. Values as (new, existing); the table
# page, retail $, installed $ and capacity (kBtu/h) are in the comments.
EIA_REFRIGERATION_COST = {
    # p. 97: 3,175 and 4,125 over 4.340
    "Commercial Reach-In Freezers": (732, 950),
    # p. 94: 2,780 and 3,643 over 2.349
    "Commercial Reach-In Refrigerators": (1183, 1551),
    # p. 91: 9,500 and 11,500 over 11.850
    "Commercial Supermarket Display Cases": (802, 970),
    # p. 104: 22,850 and 24,850 over 23.500
    "Commercial Walk-In Freezers": (972, 1057),
    # p. 100: 19,847 and 23,897 over 41.024
    "Commercial Walk-In Refrigerators": (484, 583),
    # p. 112: 2,762 and 2,762 over 4.700
    "Commercial Beverage Merchandisers": (588, 588),
    # p. 108: 3,020 over 4.580; installed is EIA's published $732/kBtu/h
    # (3,350 over 4.580 is 731.4; EIA's figure uses an unrounded capacity)
    "Commercial Ice Machines": (659, 732),
    # p. 115: 3,272 and 3,489 over 1.810
    "Commercial Refrigerated Vending Machines": (1808, 1928),
}

# Edits to specific measures requested in review, applied to the proposed
# workbook and logged in the BTB Diff sheet: old measure Name -> new Name
# (optional) and cell values for other columns. Descriptions have "120V"
# dropped from "120V HP" wherever the Name is changed this way.
MEASURE_EDITS = {
    # Refrigeration measures use EIA values entirely (cost, capacity and
    # lifetime, 2022$), not a mix of BTB and EIA: the zero-cost reference-case
    # add-on becomes a full service measure with EIA's typical retail (new)
    # and total installed (existing) costs per kBtu/h of cooling capacity.
    # (C) Best Refrigeration already holds EIA's High values.
    "(C) Ref. Case Refrigeration": {
        "_note": "EIA Appendix C typical-case values (not BTB)",
        "Measure Type": "full service",
        "Installed Cost": "; ".join(
            f"{tech}: new: {new}; {tech}: existing: {existing}"
            for tech, (new, existing) in EIA_REFRIGERATION_COST.items()),
        "Cost Source Notes": (
            "Retail (new) and total installed (existing) cost for typical "
            "efficiency refrigeration units in 2022, normalized by the "
            "typical unit's cooling capacity (kBtu/h)"),
        "Cost Source Details": (
            "Updated Buildings Sector Appliance and Equipment Costs and "
            "Efficiencies: Appendix C; U.S. Energy Information "
            "Administration; 2023; [97, 94, 91, 104, 100, 112, 108, 115]; "
            "https://www.eia.gov/analysis/studies/buildings/equipcosts/pdf/"
            "appendix-c.pdf")},
    # No 120V heat pump dryers are on the market. The fuel switch installs
    # a 240V circuit for the new load.
    "(R) Best 120V HP Dryer FS": {
        "Name": "(R) Best HP Dryer FS",
        "Electrical Upgrade Costs": "panel: false; 240V circuit: true"},
    "(R) Best 120V HP Dryer TS": {"Name": "(R) Best HP Dryer TS"},
}
_VOLTAGE_IN_DESCRIPTION = re.compile(r"120V\s?HP")

# Scout furnace token whose BTB cost is added for a dual-fuel backup.
BACKUP_FURNACE_TECH = {"gas": "furnace (NG)", "oil": "furnace (distillate)"}


def capacity_kbtuh(btb_row):
    """The BTB row's Typical capacity (regression metric 1) in kBtu/h, or
    None if metric 1 is not a capacity in a known unit. The commercial
    ranges have no cooking capacity, so their Combined Idle Energy Rate
    (BTU/h for gas, kW for electric) is used. Units BTB describes by volume
    use the capacities in BTB_UNIT_CAPACITY_KBTUH."""

    key = (btb_row.get("btb_sector"), btb_row.get("tech_id"))
    if key in BTB_UNIT_CAPACITY_KBTUH:
        return BTB_UNIT_CAPACITY_KBTUH[key]
    metric = str(btb_row.get("Regression metric 1 - Metric")).lower()
    if not any(name in metric for name in (
            "capacity", "heat output", "combined idle energy rate")):
        return None
    factor = CAPACITY_TO_KBTUH.get(
        str(btb_row.get("Regression metric 1 - Unit")).strip().lower())
    value = parse_currency(btb_row.get("Regression metric 1 - Typical"))
    if factor is None or not value:
        return None
    return value * factor


def cost_skip_reason(row):
    """Return why this row's installed cost must not be replaced with a BTB
    cost, or None if it may be."""

    units = row.get("Cost Units")
    if cost_units_info(units) is None:
        return (f"Cost Units are {units!r}, not a whole-unit or "
                "per-heating/cooling/water-heating/cooking-capacity cost; "
                "left as-is")
    return None


# "Pt." in a heat pump measure's Name marks a dual-fuel (partial) measure that
# keeps its backup furnace. Some cost notes do not say so ("also add retail
# cost of typical furnace"), so the Name is used too: "NG" in it means a gas
# backup furnace, "Other Fossil" an oil one.
_DUAL_FUEL_NAME = re.compile(r"\bPt\.")


def dual_fuel_from_name(name):
    """The backup furnace fuel ("gas"/"oil") implied by a dual-fuel heat
    pump measure's Name, or None if it is not a dual-fuel measure."""

    if not isinstance(name, str) or not _DUAL_FUEL_NAME.search(name):
        return None
    if "NG" in name:
        return "gas"
    return "oil" if "Other Fossil" in name else None


def composite_adjustment(notes, sector, crosswalk, btb_data, name=None):
    """Amounts to add to a technology's BTB cost to rebuild a composite cost
    described by the row's Cost Source Notes.

    - "also add retail cost of typical gas/oil furnace" (dual-fuel heat
      pump): the new-construction cost also includes the backup furnace's BTB
      cost; existing construction keeps the existing backup unit, so is the
      heat pump alone.
    - "tank removal": existing-construction cost includes oil-tank removal.

    Returns:
        ({"new": x, "existing": y}, note) or None if the notes describe an
        adjustment that cannot be rebuilt from BTB.
    """

    none = {"new": 0, "existing": 0}
    if not isinstance(notes, str):
        return none, ""
    match = _ALSO_ADD_FURNACE.search(notes)
    fuel = match.group(1).lower() if match else dual_fuel_from_name(name)
    if fuel:
        entry = crosswalk.get((BACKUP_FURNACE_TECH[fuel], sector, "Ref. Case"))
        btb_row = get_btb_row(btb_data, entry) if entry is not None else None
        furnace = compute_cost(btb_row)[0] if btb_row is not None else None
        if furnace is None:
            return None
        return ({"new": furnace["new"], "existing": 0},
                f"new construction includes the BTB typical {fuel} furnace "
                f"(${furnace['new']}); existing keeps the existing backup")
    if _TANK_REMOVAL.search(notes):
        return ({"new": 0, "existing": TANK_REMOVAL_COST},
                f"existing includes ${TANK_REMOVAL_COST} oil-tank removal "
                "(from the existing notes, not BTB)")
    if _UNHANDLED_COMPOSITE.search(notes):
        return None
    return none, ""


def entry_value(entry, column):
    """A crosswalk entry's value for an optional text column, or None."""

    value = entry.get(column)
    return value if isinstance(value, str) and value else None


def sector_for_name(name):
    """Infer a meas_in row's sector from the "(R)"/"(C)" prefix of its Name
    (present on every row, and consistent with its Building Type column)."""

    if isinstance(name, str):
        if name.startswith("(R)"):
            return "residential"
        if name.startswith("(C)"):
            return "commercial"
    return None


def techs_for_row(row):
    """Ordered, de-duplicated technology tokens referenced by a meas_in row."""

    seen = []
    for col in ["Baseline Technology", "Switched to Technology"]:
        val = row.get(col)
        if isinstance(val, str):
            for part in val.split(";"):
                part = part.strip()
                if part and part not in seen:
                    seen.append(part)
    return seen


def find_value_for_key(cell, key):
    """Return the raw (unparsed) value substring for `key` in a nested
    "key: value; key2: value2" cell, or the whole cell if it has no nested
    keys at all, or None if `key` is not present in a nested cell."""

    if not isinstance(cell, str):
        return None if pd.isna(cell) else str(cell)
    if ":" not in cell:
        return cell.strip()
    pattern = re.compile(
        re.escape(key) + r"\s*:\s*([^;]+)", re.IGNORECASE)
    match = pattern.search(cell)
    return match.group(1).strip() if match else None


def unit_for_tech(perf_units_cell, tech):
    """Resolve the meas_in performance unit that applies to `tech` in a
    given row, trying a direct tech-name key first and falling back to the
    row's heating/cooling/ventilation end-use key (see END_USE_FALLBACK)."""

    direct = find_value_for_key(perf_units_cell, tech)
    if direct:
        return direct
    if isinstance(perf_units_cell, str):
        for end_use, techs in END_USE_FALLBACK.items():
            if tech in techs:
                via_end_use = find_value_for_key(perf_units_cell, end_use)
                if via_end_use:
                    return via_end_use
    return None


def owner_tech(row):
    """The one technology an un-keyed cell in this row (a bare value, or a
    bare "new: X; existing: Y" cost) describes: the Switched-to technology,
    or the Baseline technology for rows with no switch (e.g. Ref. Case) or
    that switch to "same" (the technology is unchanged). None if that is not
    a single technology."""

    for col in ["Switched to Technology", "Baseline Technology"]:
        val = row.get(col)
        if isinstance(val, str) and val.strip() and val.strip() != "same":
            parts = [p.strip() for p in val.split(";") if p.strip()]
            return parts[0] if len(parts) == 1 else None
    return None


_BARE_COST = re.compile(
    r"\s*new\s*:\s*[^;:]+;\s*existing\s*:\s*[^;:]+", re.IGNORECASE)


def is_unkeyed(cell):
    """True if a meas_in cell holds a single un-keyed value (blank, numeric,
    plain text, or a bare "new: X; existing: Y" cost) rather than
    "technology: value" pairs."""

    if not isinstance(cell, str):
        return True
    return ":" not in cell or bool(_BARE_COST.fullmatch(cell))


# BTB metric names (lowercased) that describe cooling vs. heating output,
# used to pick the right "cooling:"/"heating:" key when a meas_in cell stores
# a technology's performance by end use rather than by technology name
# (e.g. "heating: 2.58; cooling: 4.4" for an ASHP).
COOLING_METRICS = {
    "seer", "seer1", "seer2", "eer", "ceer", "ieer", "cooling cop"}
HEATING_METRICS = {"hspf", "hspf2", "afue", "heating cop"}


def metric_conflicts_with_token(tech, metric_name):
    """True if `tech` is a heating-only or cooling-only token (suffix "-heat"
    or "-cool") and the crosswalk's BTB metric is for the other end use.

    BTB often gives one metric for a row that several Scout tokens map to
    (e.g. only Heating COP for ground source heat pumps), so the metric can't
    be used as the performance of the token for the other end use."""

    metric = metric_name.strip().lower() \
        if isinstance(metric_name, str) else ""
    return (tech.endswith("-cool") and metric in HEATING_METRICS) or \
        (tech.endswith("-heat") and metric in COOLING_METRICS)


def perf_key_for_tech(perf_cell, tech, metric_name, owner):
    """Return the key under which `tech`'s performance is stored in a
    meas_in Energy Performance cell: the technology name, the end-use key
    ("heating"/"cooling") implied by the BTB metric, or None if neither is
    present. A cell with no nesting holds one value, which belongs to `owner`
    (see owner_tech)."""

    if is_unkeyed(perf_cell):
        return tech if tech == owner else None
    flags = re.IGNORECASE
    if re.search(re.escape(tech) + r"\s*:", perf_cell, flags):
        return tech
    metric = metric_name.strip().lower() \
        if isinstance(metric_name, str) else ""
    end_use = "cooling" if metric in COOLING_METRICS else \
        "heating" if metric in HEATING_METRICS else None
    if end_use and re.search(end_use + r"\s*:", perf_cell, flags):
        return end_use
    return None


def replace_value_for_key(cell, key, new_value):
    """Substitute the value for `key` in a nested cell (or replace the
    whole cell if it has no nested keys), leaving all other keys/formatting
    untouched. Returns the original cell unchanged if `key` is not found in
    a nested cell (caller should not have called this in that case)."""

    if not isinstance(cell, str):
        return new_value
    if ":" not in cell:
        return new_value
    pattern = re.compile(
        r"(" + re.escape(key) + r"\s*:\s*)([^;]+)", re.IGNORECASE)
    return pattern.sub(lambda m: m.group(1) + new_value, cell, count=1)


def replace_cost_for_key(cell, key, new_cost):
    """Same as replace_value_for_key, but for the "key: new: X; key:
    existing: Y" nesting used by the Installed Cost column."""

    if is_unkeyed(cell):
        # No existing nested cost structure to preserve -- write a fresh
        # "new: X; existing: Y" pair for this (single) technology.
        return f"new: {new_cost['new']}; existing: {new_cost['existing']}"
    result = cell
    any_found = False
    for sub_key, val in new_cost.items():
        pattern = re.compile(
            r"(" + re.escape(key) + r"\s*:\s*" + re.escape(sub_key) +
            r"\s*:\s*)([^;]+)", re.IGNORECASE)
        if pattern.search(result):
            any_found = True
            result = pattern.sub(
                lambda m: m.group(1) + str(val), result, count=1)
    return result if any_found else cell


def get_btb_row(btb_data, entry):
    """Look up the single BTB row matching a crosswalk entry's technology
    ID, projection year, and projection scenario."""

    df = btb_data[entry["sector"]]
    match = df[
        (df["tech_id"] == entry["btb_technology_id"])
        & (df["Projection Year"] == entry["projection_year"])
        & (df["Projection Scenario"] == entry["projection_scenario"])]
    if len(match) != 1:
        return None
    return match.iloc[0]


def parse_currency(val):
    """Parse a BTB cost value (e.g. "$5,158" or 5158.0) into a float."""

    if pd.isna(val):
        return None
    if isinstance(val, str):
        val = val.replace("$", "").replace(",", "").strip()
        if not val:
            return None
    try:
        return float(val)
    except ValueError:
        return None


def is_zero_cost_placeholder(old_cost):
    """True if a technology's existing cost has a "new" or "existing"
    sub-value of exactly 0. A real installed cost is never $0, so this
    normally means the cost was deliberately zeroed out because the
    technology shares physical equipment with a paired heating/cooling
    technology elsewhere in the same row (e.g. a heat pump's cost entered
    once under the "heat" entry and zeroed under "cool" to avoid double-
    counting total installed cost) -- such rows are left alone rather than
    overwritten with a full BTB cost that would reintroduce that
    double-count.
    """

    if not isinstance(old_cost, str):
        return parse_currency(old_cost) == 0
    if ":" not in old_cost:
        return parse_currency(old_cost) == 0
    for part in old_cost.split(";"):
        if ":" in part:
            val = parse_currency(part.split(":", 1)[1])
            if val == 0:
                return True
    return False


def installation_multiplier(btb_row, kind):
    """Installation cost multiplier (installed = retail * multiplier, or
    retail + adder for adder-based technologies, which have multiplier 1)
    for `kind` "new" or "retrofit".

    Residential BTB rows give it directly. Commercial rows don't, so it is
    recovered from the row's Low/Mid/High retail and installed prices: per
    the BTB documentation a technology uses either an adder or a multiplier,
    so whichever of the two reproduces all three installed prices is used.

    Returns None if it cannot be determined (e.g. no retail price).
    """

    suffix = "New Construction" if kind == "new" else "Retrofit"
    direct = parse_currency(btb_row.get(f"Installation Multiplier - {suffix}"))
    if direct is not None:
        return direct
    installed_col = ("Typical New Construction Installed Cost ($2023)"
                     if kind == "new"
                     else "Typical Retrofit Installed Cost ($2023)")
    retail = [parse_currency(btb_row.get(
        f"Typical Retail Price ($2023) - {q}")) for q in ("Low", "Mid", "High")]
    installed = [parse_currency(btb_row.get(f"{installed_col} - {q}"))
                 for q in ("Low", "Mid", "High")]
    if None in retail or None in installed or retail[1] <= 0:
        return None
    adder = installed[1] - retail[1]
    if all(abs((i - r) - adder) <= 0.5 + 0.002 * abs(i)
           for r, i in zip(retail, installed)):
        return 1.0
    multiplier = installed[1] / retail[1]
    if all(abs(i - multiplier * r) <= 0.5 + 0.005 * abs(i)
           for r, i in zip(retail, installed)):
        return multiplier
    return None


def compute_cost(btb_row, metric_idx=None, perf_value=None):
    """Return {"new": ..., "existing": ...} installed cost ($2023).

    Starts from BTB's precomputed Mid installed cost, which is at the Typical
    performance level, and moves it to `perf_value` (a value of regression
    metric `metric_idx`, in BTB units) using the retail price regression:

        retail change = coef2_mid * (perf_value - typical) * unit_multiplier
        installed     = installed_typical + multiplier * retail change

    (all other regression inputs, e.g. capacity, stay at Typical). This is
    equivalent to evaluating the full regression and applying the
    installation multiplier/adder, but always agrees with BTB at Typical.
    Only regression metric 2 is treated as a performance metric.

    Returns:
        (cost dict or None, True if the cost was moved off the Typical
        performance level using the regression).
    """

    new_val = parse_currency(btb_row.get(
        "Typical New Construction Installed Cost ($2023) - Mid"))
    existing_val = parse_currency(btb_row.get(
        "Typical Retrofit Installed Cost ($2023) - Mid"))
    if new_val is None or existing_val is None:
        return None, False

    typical = parse_currency(btb_row.get("Regression metric 2 - Typical"))
    coef = parse_currency(btb_row.get("Regression metric 2 - Coefficient-Mid"))
    if metric_idx != 2 or perf_value is None or typical is None or \
            coef is None:
        return {"new": round(new_val), "existing": round(existing_val)}, False

    unit_mult = parse_currency(btb_row.get("Typical unit multiplier")) or 1.0
    retail_change = coef * (perf_value - typical) * unit_mult
    if retail_change == 0:
        return {"new": round(new_val), "existing": round(existing_val)}, True
    mult_new = installation_multiplier(btb_row, "new")
    mult_ret = installation_multiplier(btb_row, "retrofit")
    if mult_new is None or mult_ret is None:
        return {"new": round(new_val), "existing": round(existing_val)}, False
    return {"new": round(new_val + mult_new * retail_change),
            "existing": round(existing_val + mult_ret * retail_change)}, True


def residential_ashp_cooling(meas_in_df):
    """The cooling COP of the residential heat pump measures for each tier,
    {tier: COP}, taken from the first residential measure switching to ASHP
    in that tier. Commercial rows that use residential system costs are
    costed at these performance levels."""

    out = {}
    for _, row in meas_in_df.iterrows():
        name = str(row.get("Name"))
        tier = tier_for_name(name)
        if not name.startswith("(R)") or tier is None or tier in out or \
                str(row.get("Switched to Technology")).strip() != "ASHP":
            continue
        value = find_value_for_key(row.get("Energy Performance"), "cooling")
        number = parse_currency(value)
        if number:
            out[tier] = number
    return out


def apply_measure_edits(row, edits, ws, excel_row, col_index, diff_rows):
    """Apply the MEASURE_EDITS for one measure to the workbook, log each in
    `diff_rows`, and return the row with the edited values."""

    row = row.copy()
    edits = dict(edits)
    note = edits.pop("_note", "edit requested in review (not from BTB)")
    if "Name" in edits and isinstance(row.get("Description"), str):
        edits["Description"] = _VOLTAGE_IN_DESCRIPTION.sub(
            "HP", row["Description"])
    for column, new_value in edits.items():
        old_value = row.get(column)
        if old_value == new_value:
            continue
        ws.cell(row=excel_row, column=col_index[column]).value = new_value
        diff_rows.append({
            "Name": row.get("Name"), "technology": "", "column": column,
            "old_value": old_value, "new_value": new_value,
            "notes": note})
        row[column] = new_value
    return row


def main():
    """Build the BTB-proposed workbook and diff sheet."""

    btb_data = load_btb_data()
    crosswalk = pd.read_csv(CROSSWALK_PATH)
    crosswalk = crosswalk[~crosswalk["needs_review"].astype(bool)]
    crosswalk_by_tech_tier = {
        (row["scout_technology"], row["sector"], row["tier"]): row
        for _, row in crosswalk.iterrows()}

    meas_in_df = pd.read_excel(XLSX_PATH, sheet_name=SHEET_NAME)
    col_index = {name: i + 1 for i, name in enumerate(meas_in_df.columns)}

    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb[SHEET_NAME]

    res_hp_cooling = residential_ashp_cooling(meas_in_df)

    diff_rows = []
    for excel_row, (_, row) in enumerate(meas_in_df.iterrows(), start=2):
        if _BREAKTHROUGH_NAME.search(str(row.get("Name"))):
            continue
        edits = MEASURE_EDITS.get(row.get("Name"))
        if edits:
            row = apply_measure_edits(
                row, edits, ws, excel_row, col_index, diff_rows)
        tier = tier_for_name(row.get("Name"))
        if tier is None:
            continue
        # Multiple technologies in this row can share the same Energy
        # Performance/Installed Cost/Lifetime cell (e.g. a combined
        # heating+cooling package). Track each cell's contents here and
        # chain edits through it -- writing straight from the original
        # pandas row on every technology would make each technology's
        # edit clobber the previous one's, since they'd all overwrite the
        # same cell starting from its pre-edit text.
        cell_state = {
            "Energy Performance": row.get("Energy Performance"),
            "Installed Cost": row.get("Installed Cost"),
            "Lifetime": row.get("Lifetime"),
        }
        owner = owner_tech(row)
        type_logged = False
        row_sector = sector_for_name(row.get("Name"))
        row_techs = techs_for_row(row)
        row_entries = {
            t: crosswalk_by_tech_tier.get((t, row_sector, tier))
            for t in row_techs}
        if owner is None and is_unkeyed(cell_state["Installed Cost"]) \
                and is_zero_cost_placeholder(cell_state["Installed Cost"]) \
                and any(e is not None and entry_value(e, "measure_type")
                        for e in row_entries.values()):
            # A zero-cost add-on covering several technologies becomes a
            # full service measure with a cost entry per technology.
            cell_state["Installed Cost"] = "; ".join(
                f"{t}: new: 0; {t}: existing: 0" for t in row_techs
                if row_entries[t] is not None)
        cost_keys = [t for t in row_techs if isinstance(
            cell_state["Installed Cost"], str) and re.search(
                re.escape(t) + r"\s*:", cell_state["Installed Cost"])]
        # A keyed cost cell in an earlier dollar year is only updated when
        # every technology in it has a crosswalk entry, so that all its
        # costs can move to $2023 together (units are relabeled below).
        all_keys_mapped = bool(cost_keys) and all(
            row_entries.get(t) is not None for t in cost_keys)
        cost_written = set()
        row_relabeled = False
        skip_cost_reason = cost_skip_reason(row)
        cost_notes = row.get("Cost Source Notes")
        for tech in row_techs:
            # Tokens shared by residential and commercial measures (e.g.
            # "HPWH") have a crosswalk entry per sector. A token with only
            # the other sector's entry is left alone, unless it is
            # deliberately crosswalked across sectors.
            entry = crosswalk_by_tech_tier.get((tech, row_sector, tier))
            if entry is None and tech in CROSS_SECTOR_TECHS:
                entry = next(
                    (crosswalk_by_tech_tier[(tech, sector, tier)]
                     for sector in ("residential", "commercial")
                     if (tech, sector, tier) in crosswalk_by_tech_tier),
                    None)
            cost_only = False
            if entry is None and tech == "ASHP" and \
                    row_sector == "commercial" and isinstance(
                        cost_notes, str) and \
                    _RESIDENTIAL_SYSTEM_NOTE.search(cost_notes):
                entry = crosswalk_by_tech_tier.get(("ASHP", "residential", tier))
                cost_only = entry is not None
            if entry is None and tech == "all":
                # Measures with no technology token (baseline "all") are
                # crosswalked by measure name.
                entry = crosswalk_by_tech_tier.get(
                    (row.get("Name"), row_sector, tier))
            if entry is None:
                continue
            if tech in RESIDENTIAL_FURNACE_ALIAS and isinstance(
                    cost_notes, str) and _RESIDENTIAL_FURNACE_NOTE.search(
                        cost_notes):
                entry = crosswalk_by_tech_tier.get(
                    (RESIDENTIAL_FURNACE_ALIAS[tech], "residential", tier),
                    entry)
            if tech == owner == "resistance heat" and isinstance(
                    cost_notes, str) and _ELEC_BOILER_COST.search(cost_notes):
                # "resistance heat" here is the commercial electric boiler
                # named in the cost notes, not a resistance furnace.
                entry = crosswalk_by_tech_tier.get(
                    ("elec_boiler", "commercial", tier), entry)
            btb_row = get_btb_row(btb_data, entry)
            if btb_row is None:
                continue
            # In electric boiler rows the baseline boiler's entries hold
            # the values of the electric boiler switched to: take the cost
            # from the electric boiler's BTB row (below) and leave
            # performance and lifetime alone.
            holds_switched = tech in BOILER_COST_KEYS and isinstance(
                cost_notes, str) and bool(_ELEC_BOILER_COST.search(cost_notes))

            # -- Performance --
            # Ref. Case and Best take performance from BTB (Typical / High).
            # Min. Efficiency and ESTAR ("Existing" bound) keep the
            # standard-defined performance already in meas_in; it is only
            # used below as the point at which cost is evaluated.
            perf_cell = cell_state["Energy Performance"]
            units_cell = row.get("Performance Units")
            metric_idx = None
            for i in (1, 2):
                if btb_row.get(f"Regression metric {i} - Metric") == \
                        entry["btb_metric_name"]:
                    metric_idx = i
                    break
            perf_key = None if (holds_switched or cost_only) else \
                perf_key_for_tech(
                perf_cell, tech, entry["btb_metric_name"], owner)
            metric_conflict = metric_conflicts_with_token(
                tech, entry["btb_metric_name"])
            target_unit = (
                find_value_for_key(units_cell, perf_key) if perf_key
                else None) or unit_for_tech(units_cell, tech)
            old_val = find_value_for_key(perf_cell, perf_key) \
                if perf_key else None
            keep_existing = entry["bound"] == EXISTING_BOUND
            cost_perf = None  # performance (BTB units) to evaluate cost at
            if metric_idx and not keep_existing:
                raw_val = btb_row.get(
                    f"Regression metric {metric_idx} - "
                    f"{REGRESSION_BOUND_COL[entry['bound']]}")
                cost_perf = parse_currency(raw_val)
                converted = convert(
                    entry["btb_metric_name"], target_unit, raw_val) \
                    if target_unit and pd.notna(raw_val) else None
                if metric_conflict and perf_key:
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Energy Performance",
                        "old_value": old_val, "new_value": "(skipped)",
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                        "notes": f"BTB metric {entry['btb_metric_name']!r} "
                                 "is for the other end use than this token; "
                                 "performance left as-is",
                    })
                elif converted is not None and perf_key:
                    new_val_str = f"{converted:.3g}"
                    if old_val != new_val_str:
                        cell_state["Energy Performance"] = \
                            replace_value_for_key(
                                perf_cell, perf_key, new_val_str)
                        ws.cell(
                            row=excel_row,
                            column=col_index["Energy Performance"]).value \
                            = cell_state["Energy Performance"]
                        diff_rows.append({
                            "Name": row.get("Name"), "technology": tech,
                            "column": "Energy Performance",
                            "old_value": old_val, "new_value": new_val_str,
                            "btb_technology_id": entry["btb_technology_id"],
                            "btb_display_name": entry["btb_display_name"],
                            "projection_scenario":
                                entry["projection_scenario"],
                            "projection_year": entry["projection_year"],
                        })
            elif metric_idx and keep_existing and cost_only and \
                    tier in res_hp_cooling:
                cost_perf = to_btb_metric(
                    entry["btb_metric_name"], "COP", res_hp_cooling[tier],
                    parse_currency(btb_row.get(
                        f"Regression metric {metric_idx} - Typical")))
            elif metric_idx and keep_existing and target_unit and old_val \
                    and not metric_conflict:
                cost_perf = to_btb_metric(
                    entry["btb_metric_name"], target_unit, old_val,
                    parse_currency(btb_row.get(
                        f"Regression metric {metric_idx} - Typical")))

            # -- Installed cost --
            units_info = cost_units_info(row.get("Cost Units"))
            cost_btb_row, cost_metric_idx, cost_perf_value = \
                btb_row, metric_idx, cost_perf
            if entry_value(entry, "cost_at") == "Typical" and metric_idx:
                cost_perf_value = parse_currency(btb_row.get(
                    f"Regression metric {metric_idx} - Typical"))
            if holds_switched:
                boiler_entry = crosswalk_by_tech_tier.get(
                    ("elec_boiler", "commercial", tier))
                boiler_row = get_btb_row(btb_data, boiler_entry) \
                    if boiler_entry is not None else None
                if boiler_row is not None:
                    cost_btb_row, cost_metric_idx, cost_perf_value = \
                        boiler_row, None, None
            cost, used_regression = compute_cost(
                cost_btb_row, cost_metric_idx, cost_perf_value)
            cost_note = "" if used_regression else (
                "cost is BTB's Mid installed cost at Typical performance")
            if entry_value(entry, "cost_at") == "Typical":
                cost_note = (cost_note + "; " if cost_note else "") + \
                    (entry_value(entry, "notes") or
                     "cost kept at BTB Typical for every tier")

            cost_cell = cell_state["Installed Cost"]
            cost_unkeyed = is_unkeyed(cost_cell)
            cost_applies = (tech == owner) if cost_unkeyed else bool(
                re.search(re.escape(tech) + r"\s*:", cost_cell,
                          re.IGNORECASE))
            old_cost = None
            if cost_applies:
                old_cost = cost_cell if cost_unkeyed \
                    else find_value_for_key(cost_cell, tech)
            new_type = entry_value(entry, "measure_type")
            # A $0 cost is a deliberate placeholder (shared equipment, or a
            # reference-case add-on), unless the crosswalk turns the add-on
            # into a full service measure carrying the BTB cost.
            zero_placeholder = cost_applies and not new_type and \
                is_zero_cost_placeholder(old_cost)

            skip_reason = None
            if cost is not None and cost_applies and not zero_placeholder:
                skip_reason = skip_cost_reason
                adjustment = composite_adjustment(
                    cost_notes, entry["sector"], crosswalk_by_tech_tier,
                    btb_data, row.get("Name"))
                if skip_reason is None and cost_btb_row is btb_row and \
                        not used_regression and \
                        entry["bound"] != "Typical" and \
                        cost_perf_value is not None:
                    # Without a usable regression the cost is only available
                    # at Typical performance, wrong for other tiers.
                    skip_reason = (
                        "BTB cost regression not usable for this row, so "
                        "cost cannot be evaluated at this tier's "
                        "performance; left as-is")
                elif skip_reason is None and cost_btb_row is btb_row and \
                        not used_regression and \
                        entry["bound"] == EXISTING_BOUND:
                    skip_reason = (
                        "kept performance could not be placed on BTB's "
                        "metric scale, so cost cannot be evaluated at it; "
                        "left as-is")
                elif skip_reason is None and adjustment is None:
                    skip_reason = (
                        "cost notes describe an adjustment that cannot be "
                        "rebuilt from BTB; left as-is")
                elif skip_reason is None and units_info and units_info[1]:
                    capacity = capacity_kbtuh(cost_btb_row)
                    if capacity is None:
                        skip_reason = (
                            "Cost Units are per kBtu/h but BTB capacity is "
                            "not a known heating/cooling capacity unit; "
                            "left as-is")
                    else:
                        # Adders (a backup furnace, tank removal) are
                        # whole-unit costs, so they are divided by the same
                        # capacity.
                        extra, extra_note = adjustment
                        cost = {k: round((v + extra[k]) / capacity, 2)
                                for k, v in cost.items()}
                        cost_note = (cost_note + "; " if cost_note else "") \
                            + f"divided by {capacity:.4g} kBtu/h"
                        if extra_note:
                            cost_note += "; " + extra_note
                elif skip_reason is None:
                    extra, extra_note = adjustment
                    cost = {"new": cost["new"] + extra["new"],
                            "existing": cost["existing"] + extra["existing"]}
                    if extra_note:
                        cost_note = (cost_note + "; " if cost_note else "") \
                            + extra_note
                if skip_reason is None and units_info and \
                        units_info[0] != 2023 and not cost_unkeyed and \
                        not all_keys_mapped:
                    skip_reason = (
                        f"Cost Units are {row.get('Cost Units')!r} but this "
                        "row has several technology costs; replacing one "
                        "with a $2023 BTB cost would mix dollar years; "
                        "left as-is")
            if cost_applies and (zero_placeholder or skip_reason):
                diff_rows.append({
                    "Name": row.get("Name"), "technology": tech,
                    "column": "Installed Cost", "old_value": old_cost,
                    "new_value": "(skipped)",
                    "btb_technology_id": entry["btb_technology_id"],
                    "btb_display_name": entry["btb_display_name"],
                    "projection_scenario": entry["projection_scenario"],
                    "projection_year": entry["projection_year"],
                    "notes": skip_reason if skip_reason else (
                        "existing cost is $0 -- a deliberate placeholder "
                        "(e.g. cost shared with a paired technology in "
                        "this row); left as-is"),
                })
            elif cost is not None and cost_applies:
                cost_written.add(tech)
                new_full_cost_cell = replace_cost_for_key(
                    cost_cell, tech, cost)
                new_cost_repr = (
                    f"new: {cost['new']}; existing: {cost['existing']}")
                same_values = cost_unkeyed and isinstance(
                    cost_cell, str) and [
                        float(v) for v in re.findall(
                            r"-?\d+\.?\d*", cost_cell)] == [
                        float(cost["new"]), float(cost["existing"])]
                if new_full_cost_cell != cost_cell and not same_values:
                    cell_state["Installed Cost"] = new_full_cost_cell
                    ws.cell(row=excel_row,
                            column=col_index["Installed Cost"]).value = \
                        new_full_cost_cell
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Installed Cost", "old_value": old_cost,
                        "new_value": new_cost_repr,
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                        "notes": cost_note,
                    })
                old_units = row.get("Cost Units")
                if units_info and units_info[0] != 2023 and cost_unkeyed \
                        and not row_relabeled:
                    row_relabeled = True
                    new_units = relabel_cost_units(old_units)
                    ws.cell(row=excel_row,
                            column=col_index["Cost Units"]).value = new_units
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Cost Units", "old_value": old_units,
                        "new_value": new_units,
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                        "notes": "BTB costs are $2023",
                    })
                old_type = row.get("Measure Type")
                if new_type and old_type != new_type and \
                        not type_logged:
                    type_logged = True
                    ws.cell(row=excel_row,
                            column=col_index["Measure Type"]).value = new_type
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Measure Type", "old_value": old_type,
                        "new_value": new_type,
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                        "notes": "zero-cost add-on becomes a full service "
                                 "measure carrying the BTB cost",
                    })

            # -- Lifetime --
            lifetime = btb_row.get("Lifetime (Years)")
            lifetime_override = entry.get("lifetime_value")
            if pd.notna(lifetime_override):
                lifetime = float(lifetime_override)
            if pd.notna(lifetime) and not holds_switched and not cost_only:
                lifetime_cell = cell_state["Lifetime"]
                life_unkeyed = is_unkeyed(lifetime_cell)
                old_lifetime = find_value_for_key(lifetime_cell, tech)
                new_lifetime_str = f"{float(lifetime):.3g}"
                life_applies = (tech == owner) if life_unkeyed else bool(
                    re.search(re.escape(tech) + r"\s*:", lifetime_cell,
                              re.IGNORECASE))
                if life_applies and old_lifetime != new_lifetime_str:
                    keep_lifetime = entry_value(entry, "keep_lifetime") == "yes"
                    if not keep_lifetime:
                        cell_state["Lifetime"] = replace_value_for_key(
                            lifetime_cell, tech, new_lifetime_str)
                        ws.cell(row=excel_row,
                                column=col_index["Lifetime"]).value = \
                            cell_state["Lifetime"]
                    diff_rows.append({
                        "Name": row.get("Name"), "technology": tech,
                        "column": "Lifetime",
                        "old_value": old_lifetime,
                        "new_value": "(skipped)" if keep_lifetime
                        else new_lifetime_str,
                        "btb_technology_id": entry["btb_technology_id"],
                        "btb_display_name": entry["btb_display_name"],
                        "projection_scenario": entry["projection_scenario"],
                        "projection_year": entry["projection_year"],
                        "notes": (
                            f"BTB lifetime {btb_row.get('Lifetime (Years)')}"
                            " not used. "
                            + (entry_value(entry, "notes") or ""))
                        if keep_lifetime else (
                            (entry_value(entry, "notes") or "")
                            if pd.notna(lifetime_override) else ""),
                    })

        old_units = row.get("Cost Units")
        units_info = cost_units_info(old_units)
        if cost_keys and cost_written and not row_relabeled and \
                units_info and units_info[0] != 2023 and all_keys_mapped:
            new_units = relabel_cost_units(old_units)
            ws.cell(row=excel_row,
                    column=col_index["Cost Units"]).value = new_units
            unwritten = [t for t in cost_keys if t not in cost_written]
            diff_rows.append({
                "Name": row.get("Name"), "technology": "",
                "column": "Cost Units", "old_value": old_units,
                "new_value": new_units,
                "notes": "BTB costs are $2023" + (
                    f"; WARNING: {unwritten} not updated, so the row mixes "
                    "dollar years" if unwritten else "")})

    if "BTB Diff" in wb.sheetnames:
        del wb["BTB Diff"]
    diff_ws = wb.create_sheet("BTB Diff")
    diff_columns = [
        "Name", "technology", "column", "old_value", "new_value",
        "btb_technology_id", "btb_display_name", "projection_scenario",
        "projection_year", "notes"]
    diff_df = pd.DataFrame(diff_rows, columns=diff_columns).fillna("")
    diff_ws.append(diff_columns)
    for _, diff_row in diff_df.iterrows():
        diff_ws.append(list(diff_row))

    wb.save(OUT_PATH)
    print(f"Proposed {len(diff_rows)} cell changes across "
          f"{diff_df['Name'].nunique() if len(diff_df) else 0} measures.")
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
